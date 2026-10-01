"""Build a traceable, complete application-source listing (PDF, DOCX, TXT).

Blank-only lines are omitted from the display, never code or comments. Long
lines wrap visually. The JSON/CSV map retains every displayed source line's
original file and line number; the UTF-8 TXT retains original file content.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import textwrap
from pathlib import Path


def build(root, scope_path, out, mono_font, chinese_font):
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    scope = json.loads(scope_path.read_text(encoding="utf-8"))
    out.mkdir(parents=True, exist_ok=True)
    files, rows, original = [], [], []
    for file_id, name in enumerate(scope["files"], 1):
        path = root / name
        if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()):
            raise ValueError(f"Source must be a repository file: {name}")
        data = path.read_bytes()
        text = data.decode("utf-8")
        lines = text.splitlines()
        start = len(rows)
        for line_number, line in enumerate(lines, 1):
            if line.strip():
                # Presentation only: indentation and original text remain in TXT.
                segments = textwrap.wrap(line.expandtabs(4), width=99, replace_whitespace=False,
                                         drop_whitespace=False, break_on_hyphens=False) or [""]
                rows.append({"file": name, "file_id": file_id, "line": line_number,
                             "text": line, "display": segments})
        files.append({"file": name, "file_id": file_id, "sha256": hashlib.sha256(data).hexdigest(),
                      "physical_lines": len(lines), "nonblank_lines": len(rows) - start,
                      "first_listing_line": start + 1, "last_listing_line": len(rows)})
        original.append(f"# ===== FILE: {name} =====\n{text.rstrip()}\n")
    full_pages = [rows[i:i + 50] for i in range(0, len(rows), 50)]
    selection = list(range(len(full_pages))) if len(full_pages) <= 60 else (
        list(range(30)) + list(range(len(full_pages) - 30, len(full_pages))))
    manifest = {"software_name": scope["software_name"], "version": scope["version"],
                "scope": scope["scope"], "files": files, "physical_lines": sum(f["physical_lines"] for f in files),
                "listing_nonblank_lines": len(rows), "full_pages": len(full_pages),
                "submission_pages": len(selection), "selected_original_pages": [p + 1 for p in selection],
                "rule": "50 nonblank original source lines per page, except the final page. "
                        "Long lines wrap visually. If at most 60 pages, submit all; otherwise first/last 30.",
                "rule_url": scope["rule_url"], "ownership_note": "A traceable scope listing, not proof of ownership."}
    (out / "source_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    (out / "EgoHand3D_V1.0_source_full.txt").write_text("\n".join(original), encoding="utf-8")
    with (out / "source_line_map.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["full_page", "row_on_page", "listing_line", "file_id", "file", "source_line"])
        for i, row in enumerate(rows):
            writer.writerow([i // 50 + 1, i % 50 + 1, i + 1, row["file_id"], row["file"], row["line"]])

    pdfmetrics.registerFont(TTFont("SourceMono", str(mono_font)))
    pdfmetrics.registerFont(TTFont("SourceChinese", str(chinese_font)))
    width, height = A4
    title = scope["software_name"] + " " + scope["version"]
    source_pdf = out / "EgoHand3D_V1.0_source_full.pdf"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(20), Mm(18)
    section.left_margin, section.right_margin = Mm(16), Mm(16)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(9)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans Mono"
    normal.font.size = Pt(7.7)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = False
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    header = section.header.paragraphs[0]
    header.text = title
    header.runs[0].font.name = "宋体"
    header.runs[0].font.size = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.text = "EgoHand3D | 源程序 | 第 "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    pdf = canvas.Canvas(str(source_pdf), pagesize=A4)
    pdf.setTitle(title + " 源程序")
    pdf.setAuthor("EgoHand3D project")
    for page_index, page in enumerate(full_pages):
        if page_index:
            doc.add_page_break()
        display_rows = sum(len(r["display"]) for r in page)
        leading = min(13.3, 685.0 / max(1, display_rows))
        pdf.setFont("SourceChinese", 10)
        pdf.drawString(45, height - 32, title)
        pdf.setFont("SourceMono", 7)
        pdf.drawRightString(width - 43, height - 32, f"{page_index + 1} / {len(full_pages)}")
        pdf.line(44, height - 39, width - 43, height - 39)
        pdf.setFont("SourceMono", 6.3)
        pdf.drawString(45, height - 51, "FILE:LINE (margin) | full file mapping: source_line_map.csv")
        y = height - 70
        for row in page:
            prefix = f"{row['file_id']:02d}:{row['line']:04d}"
            pdf.setFont("SourceMono", 5.7)
            pdf.setFillColorRGB(.40, .40, .40)
            pdf.drawString(42, y, prefix)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(leading)
            p.paragraph_format.keep_together = True
            for i, part in enumerate(row["display"]):
                pdf.setFillColorRGB(0, 0, 0)
                pdf.setFont("SourceMono", 7.7)
                pdf.drawString(77, y, part)
                r = p.add_run((prefix + " " if i == 0 else " " * 8) + part)
                if i < len(row["display"]) - 1:
                    r.add_break()
                y -= leading
        pdf.setFont("SourceChinese", 7)
        pdf.drawString(45, 30, "源程序 · 仅排版省略空白行 · 长行视觉换行 · 文件与行号见配套清单")
        pdf.showPage()
    pdf.save()
    doc.save(str(out / "EgoHand3D_V1.0_source_full.docx"))
    # At the current application scope every page is required. Never pad to 60.
    if len(full_pages) <= 60:
        import shutil
        for suffix in ["pdf", "docx"]:
            shutil.copy2(out / f"EgoHand3D_V1.0_source_full.{suffix}",
                         out / f"EgoHand3D_V1.0_source_submission.{suffix}")
    else:
        from pypdf import PdfReader, PdfWriter
        reader, writer = PdfReader(source_pdf), PdfWriter()
        for i in selection:
            writer.add_page(reader.pages[i])
        with (out / "EgoHand3D_V1.0_source_submission.pdf").open("wb") as target:
            writer.write(target)
    sums = []
    for path in sorted(out.iterdir()):
        if path.is_file() and path.name != "SHA256SUMS":
            sums.append(hashlib.sha256(path.read_bytes()).hexdigest() + "  " + path.name)
    (out / "SHA256SUMS").write_text("\n".join(sums) + "\n")
    print(json.dumps({k: manifest[k] for k in ["physical_lines", "listing_nonblank_lines", "full_pages", "submission_pages"]}, indent=2))


def main():
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scope", type=Path, default=root / "registration/source_scope.json")
    parser.add_argument("--out", type=Path, default=root / "registration/source")
    parser.add_argument("--mono-font", type=Path, default=Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"))
    parser.add_argument("--chinese-font", type=Path, default=Path("/usr/share/fonts/truetype/arphic-gbsn00lp/gbsn00lp.ttf"))
    args = parser.parse_args()
    build(root, args.scope, args.out, args.mono_font, args.chinese_font)


if __name__ == "__main__":
    main()
